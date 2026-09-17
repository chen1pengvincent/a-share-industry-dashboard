import csv
import hashlib
import tempfile
import unittest
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from industry_workbench.exports import build_export, csv_cell, SHEET_NAMES
from industry_workbench.models import DataError, TAXONOMIES
from industry_workbench.periods import aggregate_period
from industry_workbench.sorting import add_name_sort_keys, filter_and_sort, name_sort_key, sort_rows


PERIOD = {"kind": "week", "key": "20260914", "start": "20260914", "end": "20260920", "as_of": "20260917",
          "endpoint": "20260917", "expected_days": 4, "available_days": 1, "trade_dates": ["20260914", "20260915", "20260916", "20260917"],
          "missing_dates": ["20260914", "20260915", "20260916"], "status": "DATA_GAP"}


def metric(value, *, status=None, reason=None):
    return {"value": value, "status": status or ("OK" if value is not None else "NA"),
            "reason_codes": [reason] if reason else [], "metric_date": "20260917", "valid_count": 4}


def row(tax="SW", level="L1", suffix="1", name="银行", flow="123456", pe="12.345"):
    return {"uid": f"{tax}:{level}:{suffix}", "code": "000" + suffix, "name": name, "taxonomy": tax, "level": level,
            "parent_uid": None, "membership_evidence_kind": "OFFICIAL_DATED", "status": "WITH_GAPS",
            "metrics": {"flow_cent": metric(flow, reason="PERIOD_DATA_GAP" if flow is None else None),
                        "net_mf_vol": metric("-3"), "pe_ttm_median": metric(pe, status="SMALL_SAMPLE"),
                        "pb_median": metric("0.123456789012345"), "pe_percentile": metric(None, reason="HISTORY_INSUFFICIENT"),
                        "official_pe": metric("9007199254740993"), "official_pb": metric("0"),
                        "return_5d": metric("12.34567890123456")},
            "counts": {"member_count": 4, "pe_valid": 4, "pb_valid": 4, "flow_expected": 3, "flow_received": 3}}


class FakeStore:
    def __init__(self):
        self.reads = []
        self.saved = {"batch_id": "batch-fixed", "as_of": "20260917", "publication_state": "PUBLISHED_WITH_GAPS",
                      "source": {"tree_sha256": "a" * 64}, "days": {"20260917": {"result": {"sha256": "b" * 64}}}}
    def manifest(self, batch):
        self.reads.append(batch)
        if batch != "batch-fixed":
            raise DataError("BATCH_NOT_FOUND")
        return deepcopy(self.saved)
    def read_json(self, ref):
        return {"trade_date": "20260917", "audit": {"base_complete": True, "traded_count": 5000,
                "classifications": {"CI": {"membership_complete": False, "unclassified_flow_cent": "9007199254740993"}}},
                "source_refs": [{"sha256": "c" * 64, "request": {"api_name": "moneyflow"}}]}


class FakeQuery:
    def __init__(self):
        self.store = FakeStore()
        self.calls = []
        self.rows = {(tax["id"], level["id"]): [row(tax["id"], level["id"])] for tax in TAXONOMIES for level in tax["levels"]}
    def industries(self, batch, **params):
        self.calls.append((batch, params))
        return {"batch_id": batch, "as_of": "20260917", "period": deepcopy(PERIOD),
                "rows": add_name_sort_keys(deepcopy(self.rows[(params["taxonomy"], params["level_or_series"])]))}


def params(fmt="csv", **extra):
    return {"batch_id": "batch-fixed", "taxonomy": "SW", "level_or_series": "L1", "period_kind": "week", "period_key": "20260917",
            "format": fmt, "page": "fusion", "view": "detail", "valuation_basis": "median", "sort_key": None,
            "sort_direction": "default", "query": "", **extra}


class SortingTests(unittest.TestCase):
    def test_pinyin_phrase_collation_and_deterministic_ties(self):
        self.assertEqual(name_sort_key("重庆银行"), "chong qing yin hang")
        rows = [row(suffix="3", name="银行"), row(suffix="2", name="阿尔法"), row(suffix="1", name="重庆")]
        self.assertEqual([r["name"] for r in sort_rows(rows, page="valuation")], ["阿尔法", "重庆", "银行"])
        ties = [row(suffix="3", name="中国"), row(suffix="1", name="中国")]
        self.assertEqual([r["code"] for r in sort_rows(ties, "name", "desc")], ["0001", "0003"])

    def test_exact_decimal_default_reset_and_missing_last_both_directions(self):
        rows = [row(suffix="3", flow=None), row(suffix="2", flow="9007199254740993"),
                row(suffix="1", flow="9007199254740992"), row(suffix="0", flow="-1")]
        self.assertEqual([r["code"] for r in sort_rows(rows, "flow_cent", "asc")], ["0000", "0001", "0002", "0003"])
        self.assertEqual([r["code"] for r in sort_rows(rows, "flow_cent", "desc")], ["0002", "0001", "0000", "0003"])
        self.assertEqual([r["code"] for r in sort_rows(rows, "name", "default", page="moneyflow")], ["0002", "0001", "0000", "0003"])
        self.assertEqual([r["code"] for r in sort_rows(rows, page="fusion")], ["0000", "0001", "0002", "0003"])

    def test_ranking_defaults_and_filter_share_page_semantics(self):
        rows = [row(suffix="1", name="Bank"), row(suffix="2", name="Bank B")]
        rows[0]["metrics"]["return_5d"] = metric(None)
        self.assertEqual([r["code"] for r in sort_rows(rows, page="valuation", view="ranking")], ["0002", "0001"])
        self.assertEqual(len(filter_and_sort(rows, params(query=" BANK "))), 2)
        self.assertEqual(len(filter_and_sort(rows, params(query="0002"))), 1)
        with self.assertRaises(DataError):
            sort_rows(rows, "__dict__", "asc")


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.directory = Path(self.tmp.name)
        self.query = FakeQuery()

    def csv_read(self, path):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))

    def test_csv_current_page_filter_decimal_sort_precision_and_batch(self):
        self.query.rows[("SW", "L1")] = [row(suffix="1", name="银行一", flow="9007199254740992"),
                                        row(suffix="2", name="银行二", flow="9007199254740993"), row(suffix="3", name="煤炭", flow="-123")]
        path = self.directory / "资金.csv"
        result = build_export(self.query, params(page="moneyflow", query="银行", sort_key="flow_cent", sort_direction="desc"), path)
        rows = self.csv_read(path)
        self.assertEqual([r["行业代码"] for r in rows], ["0002", "0001"])
        self.assertEqual(rows[0]["净流入原值（0.01万元）"], "9007199254740993")
        self.assertEqual(rows[0]["净流入（亿元）"], "9007199254.740993")
        self.assertNotIn("PE_TTM中位数", rows[0])
        self.assertEqual(rows[0]["净流入量（手）"], "-3")
        self.assertTrue(all(r["批次"] == "batch-fixed" for r in rows))
        self.assertEqual(result["sha256"], hashlib.sha256(path.read_bytes()).hexdigest())
        self.assertEqual(len(self.query.calls), 1)
        self.assertEqual(self.query.store.reads, ["batch-fixed"])
        self.assertEqual(result["audit"]["verified_rows"], 2)

    def test_csv_formula_injection_quote_newline_and_negative_numbers(self):
        self.query.rows[("SW", "L1")] = [row(name=' =SUM(1,2),"文\n本"', flow="-123")]
        path = self.directory / "safe.csv"
        build_export(self.query, params(), path)
        record = self.csv_read(path)[0]
        self.assertEqual(record["行业名称"], "' =SUM(1,2),\"文\n本\"")
        self.assertEqual(record["净流入原值（0.01万元）"], "-123")
        for value in ("=1", "+1", "-1", "@cmd", "\tvalue", "\rvalue", "\nvalue", " \ufeff=1"):
            self.assertTrue(csv_cell(value).startswith("'"))
        self.assertEqual(csv_cell(Decimal("-1")), "-1")

    def test_csv_parent_filter_precedes_search_and_sort_xlsx_ignores_it(self):
        rows = [row(level="L2", suffix="1", name="行业甲", flow="2"), row(level="L2", suffix="2", name="行业乙", flow="3")]
        rows[0]["parent_uid"], rows[1]["parent_uid"] = "SW2021:L1:100000", "SW2021:L1:200000"
        self.query.rows[("SW", "L2")] = rows
        path = self.directory / "parent.csv"
        build_export(self.query, params(level_or_series="L2", parent_uid="SW2021:L1:100000", query="行业", sort_key="flow_cent", sort_direction="desc"), path)
        self.assertEqual([r["行业代码"] for r in self.csv_read(path)], ["0001"])
        all_path = self.directory / "parent.xlsx"
        result = build_export(self.query, params("xlsx", parent_uid="SW2021:L1:100000"), all_path)
        self.assertEqual(result["row_count"], 10)
        with self.assertRaisesRegex(DataError, "INVALID_PARENT_FILTER"):
            build_export(self.query, params(parent_uid=["wrong"]), self.directory / "bad-parent.csv")

    def test_csv_na_reason_date_and_small_sample_are_explicit(self):
        self.query.rows[("SW", "L1")] = [row(flow=None)]
        path = self.directory / "na.csv"
        build_export(self.query, params(), path)
        record = self.csv_read(path)[0]
        self.assertEqual(record["净流入（亿元）"], "NA")
        self.assertIn("PERIOD_DATA_GAP", record["净流入（亿元）：缺失原因"])
        self.assertEqual(record["PE_TTM中位数：状态"], "小样本")
        self.assertEqual(record["PE_TTM中位数：数据日期"], "20260917")
        self.assertEqual(record["PE_TTM中位数：有效样本数"], "4")

    def test_flow_coverage_diagnostics_keep_record_and_trade_day_units_distinct(self):
        daily = row()
        daily["metrics"]["flow_cent"].update(received_count=3, expected_count=3)
        day_period = {**PERIOD, "kind": "day", "key": "20260917", "start": "20260917", "end": "20260917",
                      "trade_dates": ["20260917"], "expected_days": 1, "available_days": 1, "missing_dates": [], "status": "COMPLETE"}
        missing_day = aggregate_period({}, day_period, [daily])
        day_inputs = {date: {"industries": [deepcopy(daily)]} for date in ("20260916", "20260917")}
        week = aggregate_period(day_inputs, deepcopy(PERIOD), [daily])
        month_period = {**PERIOD, "kind": "month", "key": "2026-09", "start": "20260901", "end": "20260930",
                        "trade_dates": ["20260901", "20260916", "20260917"], "expected_days": 3}
        month = aggregate_period(day_inputs, month_period, [daily])
        cases = [("day", day_period, daily, ("3", "3")),
                 ("missing-day", missing_day["period"], missing_day["industries"][0], ("NA", "NA")),
                 ("week", week["period"], week["industries"][0], ("2", "4")),
                 ("month", month["period"], month["industries"][0], ("2", "3"))]
        # The missing-day projection counts one absent trade day. That must
        # never become one expected stock record or zero known stock records.
        self.assertEqual(missing_day["industries"][0]["metrics"]["flow_cent"]["expected_count"], 1)
        for name, period, industry, expected in cases:
            with self.subTest(period=name):
                self.query.industries = lambda batch, **kwargs: {"batch_id": batch, "as_of": "20260917", "period": deepcopy(period), "rows": [deepcopy(industry)]}
                output = self.directory / (name + ".csv")
                build_export(self.query, params(page="moneyflow", period_kind=period["kind"], period_key=period["key"]), output)
                record = self.csv_read(output)[0]
                labels = ("资金已取记录数", "资金应有记录数") if period["kind"] == "day" else ("资金有效交易日数", "资金应有交易日数")
                other = ("资金有效交易日数", "资金应有交易日数") if period["kind"] == "day" else ("资金已取记录数", "资金应有记录数")
                self.assertEqual(tuple(record[label] for label in labels), expected)
                self.assertTrue(all(label not in record for label in other))

    def test_xlsx_exact_six_sheets_all_levels_and_original_precision(self):
        path = self.directory / "四分类.xlsx"
        result = build_export(self.query, params("xlsx", query="必不匹配", sort_key="name", sort_direction="asc"), path)
        book = load_workbook(path, data_only=False)
        self.addCleanup(book.close)
        self.assertEqual(tuple(book.sheetnames), SHEET_NAMES)
        self.assertEqual(result["row_count"], 9)
        self.assertEqual(len(self.query.calls), 9)
        self.assertEqual({call[0] for call in self.query.calls}, {"batch-fixed"})
        for name, levels in (("申万行业", {"L1", "L2", "L3"}), ("中信行业", {"L1", "L2", "L3"}), ("通达信行业", {"880", "881"})):
            sheet = book[name]
            self.assertEqual({sheet.cell(r, 5).value for r in range(2, sheet.max_row + 1)}, levels)
            self.assertEqual(sheet.freeze_panes, "B2")
            self.assertTrue(sheet.auto_filter.ref)
        sheet = book["申万行业"]
        headers = {c.value: c.column for c in sheet[1]}
        high = sheet.cell(2, headers["官方PE"])
        self.assertEqual(high.value, "9007199254740993")
        self.assertEqual(high.data_type, "s")
        self.assertEqual(sheet.cell(2, headers["PB中位数"]).value, 0.123456789012345)
        self.assertEqual(sheet.cell(2, headers["官方PB"]).value, 0)
        self.assertEqual(sheet.cell(2, headers["5交易日收益（%）"]).value, "12.34567890123456")
        self.assertTrue(any("moneyflow" in str(c.value) for r in book["审计"] for c in r))
        self.assertTrue(result["audit"]["precision_verified"])

    def test_xlsx_text_is_not_formula_url_or_date_coercion(self):
        self.query.rows[("SW", "L1")] = [row(name="=1+1", suffix="00001")]
        self.query.rows[("SW", "L1")][0]["metrics"]["pb_median"] = metric("0.00000000000000001")
        path = self.directory / "literal.xlsx"
        build_export(self.query, params("xlsx"), path)
        book = load_workbook(path, data_only=False)
        self.addCleanup(book.close)
        self.assertEqual(book["申万行业"]["A2"].value, "=1+1")
        self.assertEqual(book["申万行业"]["A2"].data_type, "s")
        self.assertEqual(book["申万行业"]["B2"].value, "00000001")
        headers = {c.value: c.column for c in book["申万行业"][1]}
        self.assertEqual(book["申万行业"].cell(2, headers["PB中位数"]).value, "0.00000000000000001")
        self.assertTrue(all(c.data_type not in {"f", "e"} for s in book for r in s for c in r))

    def test_empty_filtered_csv_still_contains_headers(self):
        path = self.directory / "empty.csv"
        result = build_export(self.query, params(query="没有"), path)
        self.assertEqual(result["row_count"], 0)
        self.assertEqual(self.csv_read(path), [])
        self.assertIn("行业名称", path.read_text(encoding="utf-8-sig"))

    def test_invalid_or_unpinned_export_never_publishes(self):
        path = self.directory / "bad.csv"
        for values in (params(batch_id=""), params(page="unknown"), params(view="unknown"),
                       params(valuation_basis="replacement"), params(sort_key="close", sort_direction="asc")):
            with self.subTest(values=values), self.assertRaises(DataError):
                build_export(self.query, values, path)
            self.assertFalse(path.exists())
        original = self.query.industries
        def changed(*args, **kwargs):
            response = original(*args, **kwargs)
            response["batch_id"] = "new-current"
            return response
        self.query.industries = changed
        with self.assertRaisesRegex(DataError, "EXPORT_BATCH_MISMATCH"):
            build_export(self.query, params(), path)
        self.assertFalse(path.exists())

    def test_write_failure_removes_staging_and_does_not_overwrite(self):
        path = self.directory / "saved.csv"
        path.write_text("existing")
        with self.assertRaisesRegex(DataError, "EXPORT_DESTINATION_EXISTS"):
            build_export(self.query, params(), path)
        self.assertEqual(path.read_text(), "existing")
        other = self.directory / "new.csv"
        with patch("industry_workbench.exports._write_csv", side_effect=DataError("EXPECTED_FAILURE")):
            with self.assertRaisesRegex(DataError, "EXPECTED_FAILURE"):
                build_export(self.query, params(), other)
        self.assertFalse(other.exists())
        self.assertFalse(list(self.directory.glob(".export-*")))

    def test_nonfinite_value_and_wrong_period_fail_before_publication(self):
        path = self.directory / "invalid.xlsx"
        self.query.rows[("SW", "L1")][0]["metrics"]["official_pe"]["value"] = "NaN"
        with self.assertRaisesRegex(DataError, "EXPORT_NONFINITE"):
            build_export(self.query, params("xlsx"), path)
        self.assertFalse(path.exists())
        self.query = FakeQuery()
        original = self.query.industries
        def wrong_period(*args, **kwargs):
            result = original(*args, **kwargs)
            if kwargs["taxonomy"] == "CI":
                result["period"]["key"] = "20260907"
            return result
        self.query.industries = wrong_period
        with self.assertRaisesRegex(DataError, "EXPORT_PERIOD_MISMATCH"):
            build_export(self.query, params("xlsx"), path)
        self.assertFalse(path.exists())


if __name__ == "__main__":
    unittest.main()
