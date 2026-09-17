"""Batch-pinned CSV/XLSX exports with independent saved-file verification."""
from __future__ import annotations

import csv
import hashlib
import math
import os
import re
import tempfile
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from .models import DataError, TAXONOMIES, decimal_text, decimal_value, json_bytes
from .sorting import PAGES, VIEWS, filter_and_sort, sort_rows

SHEET_NAMES = ("申万行业", "同花顺行业", "通达信行业", "中信行业", "口径说明", "审计")
EXPORT_PARAMS = {"batch_id", "taxonomy", "level_or_series", "period_kind", "period_key", "format", "query",
                 "sort_key", "sort_direction", "page", "view", "valuation_basis", "parent_uid"}
METRICS = {
    "flow_cent": "净流入（亿元）", "net_mf_vol": "净流入量（手）",
    "pe_ttm_median": "PE_TTM中位数", "pb_median": "PB中位数",
    "pe_percentile": "PE历史百分位（%）", "pb_percentile": "PB历史百分位（%）",
    "official_pe": "官方PE", "official_pb": "官方PB",
    "official_pe_percentile": "官方PE历史百分位（%）", "official_pb_percentile": "官方PB历史百分位（%）",
    "close": "官方指数收盘点位", "return_5d": "5交易日收益（%）",
    "return_mtd": "月初至今收益（%）", "return_ytd": "年初至今收益（%）",
}
COUNTS = {"member_count": "期末已确认成员数", "pe_valid": "PE有效样本数", "pb_valid": "PB有效样本数",
          "flow_expected": "期末应有资金记录数", "flow_received": "期末已取资金记录数",
          "unknown_member_count": "归属未确认成员数", "basic_received": "已取估值记录数"}
PAGE_KEYS = {
    "moneyflow": ("flow_cent", "net_mf_vol", "member_count", "flow_expected", "flow_received"),
    "valuation": ("pe_ttm_median", "pe_percentile", "pb_median", "pb_percentile", "official_pe",
                  "official_pe_percentile", "official_pb", "official_pb_percentile", "return_5d", "return_mtd",
                  "return_ytd", "member_count", "pe_valid", "pb_valid"),
    "fusion": ("flow_cent", "pe_ttm_median", "pe_percentile", "pb_median", "pb_percentile", "official_pe",
               "official_pb", "member_count", "pe_valid", "pb_valid"),
}
HUMAN = {
    "OK": "完整", "SMALL_SAMPLE": "小样本", "NA": "暂无数据", "WITH_GAPS": "存在缺口",
    "COMPLETE": "完整", "IN_PROGRESS": "期间进行中", "DATA_GAP": "期间数据缺失", "NO_TRADING_DAYS": "无交易日",
    "OFFICIAL_DATED": "官方带日期成员", "OBSERVED_SAME_DAY": "当日采集归属", "UNKNOWN": "未确认",
    "HISTORY_UNAVAILABLE": "缺少历史归属", "MEMBERSHIP_HISTORY_UNAVAILABLE": "缺少历史归属",
    "MEMBERSHIP_COVERAGE_UNKNOWN": "归属覆盖未确认", "MEMBERSHIP_BOUNDARY_UNKNOWN": "归属日期边界未确认",
    "MEMBERSHIP_OVERLAP_UNKNOWN": "归属重叠未确认", "MEMBER_LIFECYCLE_UNKNOWN": "股票生命周期未确认",
    "NO_VALID_VALUES": "无有效正值", "HISTORY_INSUFFICIENT": "历史有效样本不足",
    "PERIOD_DATA_GAP": "期间数据缺失", "PERIOD_ENDPOINT_MISSING": "期末数据缺失",
    "CURRENT_VALUE_MISSING": "当前值缺失", "RETURN_ANCHOR_MISSING": "收益基准日缺失",
    "OFFICIAL_VALUE_MISSING": "官方数值缺失", "NOT_PUBLISHED": "官方未发布", "NA_NOT_PUBLISHED": "官方未发布",
}


@dataclass
class Table:
    headers: list[str]
    rows: list[list]


def _human(value):
    return HUMAN.get(value, value) if value is not None else "NA"


def _number(value):
    if value is None:
        return None
    result = decimal_value(value)
    if result is None:
        raise DataError("EXPORT_NONFINITE_OR_INVALID_NUMBER")
    return result


def _flow_yi(value):
    number = _number(value)
    if number is None:
        return None
    sign, digits, exponent = number.as_tuple()
    return Decimal((sign, digits, exponent - 6))


def _plain(value):
    if value is None:
        return "NA"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (dict, list)):
        return json_bytes(value).decode().rstrip("\n")
    return str(value)


def csv_cell(value):
    """Protect text against spreadsheet formula execution; numbers remain signed."""
    if isinstance(value, Decimal):
        return decimal_text(value)
    text = _plain(value)
    if "\x00" in text:
        raise DataError("EXPORT_INVALID_TEXT")
    meaningful = text.lstrip(" \t\r\n\ufeff")
    if meaningful.startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def _excel_value(value):
    if isinstance(value, Decimal):
        # Excel promises 15 significant digits. Additional roundtrip check also
        # catches finite values outside its practical number range.
        digits = value.as_tuple().digits
        significant = len("".join(map(str, digits)).lstrip("0").rstrip("0")) or 1
        number = float(value)
        # The displayed fixed-point format has 15 fractional places. Store
        # finer-scale values as text as well, so a nonzero value never displays
        # as zero and visible digits are never silently rounded away.
        if significant <= 15 and value.as_tuple().exponent >= -15 and math.isfinite(number) and Decimal(str(number)) == value:
            return number, "number"
        return decimal_text(value), "text"
    return _plain(value), "text"


def _industry_table(rows, batch, period, keys, *, publication_state, include_rank=False):
    headers = ["行业名称", "行业代码", "行业唯一标识", "分类体系", "层级或系列"]
    if include_rank:
        headers += ["同层级或系列资金排名"]
    headers += [METRICS.get(k, COUNTS.get(k, k)) for k in keys]
    headers += ["行业完整性", "归属证据", "批次", "批次发布状态", "周期", "周期标识", "期间开始", "期间结束",
                "已观察至", "期末交易日", "周期状态", "应有交易日数", "已取交易日数", "缺失交易日", "父级唯一标识"]
    metric_keys = [k for k in keys if k in METRICS]
    for key in metric_keys:
        headers += [METRICS[key] + suffix for suffix in ("：状态", "：缺失原因", "：数据日期", "：有效样本数")]
    if "flow_cent" in keys:
        count_headers = ["资金已取记录数", "资金应有记录数"] if period["kind"] == "day" else ["资金有效交易日数", "资金应有交易日数"]
        headers += ["净流入原值（0.01万元）", *count_headers, "资金缺失日期", "已知资金小计（0.01万元，仅诊断）"]
    output = []
    for row in rows:
        values = [row.get("name"), row.get("code"), row.get("uid"),
                  next((t["name"] for t in TAXONOMIES if t["id"] == row.get("taxonomy")), row.get("taxonomy")), row.get("level")]
        if include_rank:
            values += [_number(row.get("flow_rank"))]
        for key in keys:
            if key in METRICS:
                raw = row.get("metrics", {}).get(key, {}).get("value")
                values.append(_flow_yi(raw) if key == "flow_cent" else _number(raw))
            else:
                values.append(_number(row.get("counts", {}).get(key)))
        values += [_human(row.get("status")), _human(row.get("membership_evidence_kind")), batch, publication_state,
                   {"day": "日", "week": "自然周", "month": "自然月"}.get(period["kind"]), period["key"],
                   period.get("start"), period.get("end"), period.get("as_of"), period.get("endpoint"), _human(period.get("status")),
                   _number(period.get("expected_days")), _number(period.get("available_days")),
                   "、".join(period.get("missing_dates", [])), row.get("parent_uid")]
        for key in metric_keys:
            m = row.get("metrics", {}).get(key, {})
            reasons = m.get("reason_codes", [])
            values += [_human(m.get("status", "NA")), "；".join(f"{_human(x)} [{x}]" for x in reasons),
                       m.get("metric_date"), _number(m.get("valid_count"))]
        if "flow_cent" in keys:
            flow = row.get("metrics", {}).get("flow_cent", {})
            counts = row.get("counts", {})
            received = counts.get("flow_received") if period["kind"] == "day" else flow.get("received_count")
            expected = counts.get("flow_expected") if period["kind"] == "day" else flow.get("expected_count")
            values += [_number(flow.get("value")), _number(received), _number(expected),
                       "、".join(flow.get("missing_dates", [])), _number(flow.get("known_subtotal"))]
        output.append(values)
    return Table(headers, output)


def _notes(batch, manifest, period, params):
    notes = [
        ("固定批次", batch), ("批次截至", manifest["as_of"]), ("批次发布状态", manifest.get("publication_state")),
        ("周期", period["kind"] + ":" + period["key"]), ("期末交易日", period.get("endpoint")),
        ("导出范围", "四分类全部层级；申万L1/L2/L3、同花顺行业、通达信880/881、中信L1/L2/L3"),
        ("当前页面口径", params["page"] + "/" + params["view"] + "/" + params["valuation_basis"]),
        ("Excel与当前筛选", "Excel固定导出同批次同周期全部四分类，不应用当前页面搜索或排序。当前视图请使用CSV。"),
        ("共同股票范围", "沪深A股，按目标日上市/退市生命周期筛选；北交所和B股不参与。"),
        ("自建估值", "按当日有效成员分别取PE_TTM、PB的正且有限值中位数；1至4个有效值标小样本，0个为NA。"),
        ("官方估值", "申万和通达信的官方PE/PB独立列示；不用于填补自建中位数。同花顺及中信无官方值时保持NA。"),
        ("历史百分位", "同一行业身份、指标与证据口径下，截至当日不晚于当日的有效值；小于或等于当前值的比例，至少252个有效交易日。"),
        ("资金口径", "个股Tushare moneyflow.net_mf_amount按每日当时行业归属聚合。原始单位万元，存储flow_cent单位0.01万元；显示亿元=flow_cent/1000000。"),
        ("资金量", "net_mf_vol单位手，正值为净流入，负值为净流出。"),
        ("自然周/月", "资金按期间内每个交易日当时归属先聚合再求和；估值、成员数及百分位取期末交易日。期间尚未结束保留进行中状态。"),
        ("成员变动", "期间资金保留退出成员历史贡献，不能用期末成员篮子回算全期资金。"),
        ("历史归属", "同花顺只使用真实同日采集成员，缺历史档案不以今日成员回填；边界或互斥归属不明的受影响指标为NA。"),
        ("跨行业合计", "不同分类不可相加；申万/中信不同层级不可相加；通达信不同系列不可相加；同花顺允许多归属，跨板块总额可能重复。"),
        ("中信覆盖", "未知归属诊断桶不参加排名。旧99.9%和10只阈值仅解释旧结果，不能使新指标从UNKNOWN变为完整。"),
        ("完整性", "原始接口失败、截断、schema异常或全市场股票行情/资金漏行使整批失败；归属或历史缺口可以带缺口发布且相关指标为NA。"),
        ("缺失值", "NA表示不可计算或无有效值，与0不同；空缺原因、日期、样本与覆盖数随值导出。已知资金小计仅诊断，不能替代完整期间资金。"),
        ("排序", "名称按pypinyin 0.55.0拼音排序，原名称及行业唯一标识为稳定次键；金额精确比较，NA升降序都置末。"),
        ("收益率", "使用官方指数收盘点位；5交易日收益与第五个此前交易日比较，MTD/YTD与月/年开始前最后交易日比较。百分数字段已是百分数，不再乘100。"),
        ("Excel精度", "超过15位有效数字或不能安全读回的金融数值使用文本保存全部原值；其余金融数值为数值单元格，零值显示0。"),
        ("来源", "Tushare官方接口；具体原始响应SHA-256、接口及日期见审计。"),
        ("用途", "研究展示；无交易、下单或对外发布权限。"),
    ]
    return Table(["口径项目", "说明"], [[key, value] for key, value in notes])


def _audit(query, manifest, period, rows_by_tax):
    output = [["批次", manifest["as_of"], manifest["batch_id"], "manifest SHA-256", hashlib.sha256(json_bytes(manifest)).hexdigest()],
              ["源码", manifest["as_of"], manifest["batch_id"], "源码树SHA-256", manifest.get("source", {}).get("tree_sha256")],
              ["完整性", period.get("as_of"), manifest["batch_id"], "周期状态", _human(period.get("status"))]]
    for tax, rows in rows_by_tax.items():
        output.append(["分类", period.get("as_of"), tax, "导出行业数", Decimal(len(rows))])
        for row in rows:
            for key, m in row.get("metrics", {}).items():
                if m.get("value") is None or m.get("reason_codes"):
                    output.append(["指标缺口", m.get("metric_date"), row["uid"], METRICS.get(key, key),
                                   "；".join(m.get("reason_codes", [])) or "NA"])

    def flatten(prefix, value):
        if isinstance(value, dict):
            for key in sorted(value):
                # Audit schema is token-free; still never copy credential-like
                # fields from unexpected metadata into a user-facing workbook.
                if any(word in key.lower() for word in ("token", "secret", "authorization", "password")):
                    raise DataError("EXPORT_SENSITIVE_AUDIT_FIELD")
                yield from flatten(f"{prefix}.{key}" if prefix else key, value[key])
        else:
            text = _plain(value)
            if len(text) <= 30000:
                yield prefix, text
            elif isinstance(value, list):
                for i, item in enumerate(value):
                    yield from flatten(f"{prefix}[{i}]", item)
            else:
                raise DataError("EXPORT_AUDIT_CELL_TOO_LONG")

    for day in period.get("trade_dates", []):
        ref = manifest.get("days", {}).get(day, {}).get("result")
        if ref is None:
            output.append(["缺失交易日", day, manifest["batch_id"], "日快照", "未采集"])
            continue
        result = query.store.read_json(ref)
        if result.get("trade_date") != day:
            raise DataError("EXPORT_DAY_IDENTITY_MISMATCH")
        output.append(["日快照", day, manifest["batch_id"], "结果SHA-256", ref.get("sha256")])
        for key, value in flatten("", result.get("audit", {})):
            output.append(["日审计", day, manifest["batch_id"], key, value])
        for source in result.get("source_refs", []):
            request = source.get("request", {})
            output.append(["原始响应", day, request.get("api_name", request.get("api")), "响应SHA-256", source.get("sha256", request.get("raw_sha256"))])
    return Table(["审计类型", "日期", "对象", "审计项", "值或结果"], output)


def _write_csv(path, table):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\r\n")
        writer.writerow(table.headers)
        writer.writerows([[csv_cell(v) for v in row] for row in table.rows])
    with path.open(encoding="utf-8-sig", newline="") as stream:
        saved = list(csv.reader(stream))
    expected = [table.headers] + [[csv_cell(v) for v in row] for row in table.rows]
    if saved != expected:
        raise DataError("EXPORT_CSV_READBACK_MISMATCH")
    return {"reader": "python.csv", "verified_rows": len(table.rows), "formula_safe": True}


def _write_xlsx(path, tables):
    import xlsxwriter
    from openpyxl import load_workbook

    workbook = xlsxwriter.Workbook(str(path), {"strings_to_formulas": False, "strings_to_urls": False})
    header = workbook.add_format({"bold": True, "bg_color": "#17365D", "font_color": "#FFFFFF", "text_wrap": True, "valign": "vcenter"})
    text_format = workbook.add_format({"num_format": "@", "valign": "top"})
    number_format = workbook.add_format({"num_format": "0.###############;[Red]-0.###############;0", "valign": "top"})
    wrap_format = workbook.add_format({"text_wrap": True, "valign": "top"})
    expected = {}
    try:
        for name, table in tables.items():
            if len(table.rows) >= 1048576 or len(table.headers) > 16384:
                raise DataError("EXPORT_EXCEL_SIZE_LIMIT")
            sheet = workbook.add_worksheet(name)
            sheet.freeze_panes(1, 1)
            sheet.autofilter(0, 0, len(table.rows), len(table.headers) - 1)
            sheet.hide_gridlines(2)
            sheet.set_row(0, 46)
            sheet.set_column(0, len(table.headers) - 1, 19)
            sheet.set_column(0, 0, 24)
            if name == "口径说明":
                sheet.set_column(1, 1, 108, wrap_format)
            elif name == "审计":
                sheet.set_column(2, 3, 38)
                sheet.set_column(4, 4, 80, wrap_format)
            all_rows = [table.headers] + table.rows
            expected[name] = []
            for ri, row in enumerate(all_rows):
                saved_row = []
                for ci, raw in enumerate(row):
                    value, kind = _excel_value(raw)
                    if kind == "text" and (len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)):
                        raise DataError("EXPORT_INVALID_EXCEL_TEXT")
                    fmt = header if ri == 0 else number_format if kind == "number" else wrap_format if name in ("口径说明", "审计") else text_format
                    result = sheet.write_number(ri, ci, value, fmt) if kind == "number" else sheet.write_string(ri, ci, value, fmt)
                    if result != 0:
                        raise DataError("EXPORT_EXCEL_WRITE_FAILED")
                    saved_row.append((value, kind, raw))
                expected[name].append(saved_row)
    finally:
        workbook.close()
    book = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    verified_cells = 0
    try:
        if tuple(book.sheetnames) != SHEET_NAMES or getattr(book, "_external_links", []):
            raise DataError("EXPORT_EXCEL_STRUCTURE_MISMATCH")
        for name, rows in expected.items():
            sheet = book[name]
            if sheet.max_row != len(rows) or sheet.max_column != len(rows[0]) or sheet.freeze_panes != "B2" or not sheet.auto_filter.ref:
                raise DataError("EXPORT_EXCEL_STRUCTURE_MISMATCH")
            for ri, row in enumerate(rows, 1):
                for ci, (value, kind, raw) in enumerate(row, 1):
                    cell = sheet.cell(ri, ci)
                    # Blank strings may read back as None in OOXML; they mean
                    # no reason/no missing dates, while numeric absence is NA.
                    actual = "" if cell.value is None and value == "" else cell.value
                    if cell.data_type in {"f", "e"} or cell.hyperlink is not None or actual != value:
                        raise DataError("EXPORT_EXCEL_READBACK_MISMATCH")
                    if isinstance(raw, Decimal) and decimal_value(actual) != raw:
                        raise DataError("EXPORT_EXCEL_PRECISION_LOSS")
                    if kind == "text" and actual != "" and cell.data_type != "s":
                        raise DataError("EXPORT_EXCEL_TYPE_MISMATCH")
                    verified_cells += 1
    finally:
        book.close()
    return {"reader": "openpyxl", "verified_cells": verified_cells, "sheets": list(SHEET_NAMES), "formula_safe": True, "precision_verified": True}


def build_export(query, params: dict, destination: Path) -> dict:
    """Build one new file from a specified immutable batch, never a latest pointer."""
    params = dict(params)
    if set(params) - EXPORT_PARAMS:
        raise DataError("INVALID_EXPORT_PARAMETER")
    params.setdefault("page", "fusion")
    params.setdefault("view", "detail")
    params.setdefault("valuation_basis", "median")
    if params["page"] not in PAGES or params["view"] not in VIEWS or params["valuation_basis"] not in {"median", "official"}:
        raise DataError("INVALID_EXPORT_VIEW")
    fmt = params.get("format")
    destination = Path(destination)
    if fmt not in {"csv", "xlsx"} or destination.suffix.lower() != "." + fmt:
        raise DataError("INVALID_EXPORT_FORMAT")
    if destination.exists():
        raise DataError("EXPORT_DESTINATION_EXISTS")
    batch = params.get("batch_id")
    if not isinstance(batch, str) or not batch:
        raise DataError("EXPORT_BATCH_REQUIRED")
    manifest = query.store.manifest(batch)
    if manifest.get("batch_id") != batch:
        raise DataError("EXPORT_BATCH_MISMATCH")
    common = {"period_kind": params.get("period_kind", "day"), "period_key": params.get("period_key")}
    period = None
    rows_by_tax = {}

    def fetch(taxonomy, level):
        nonlocal period
        response = query.industries(batch, taxonomy=taxonomy, level_or_series=level, **common)
        if response.get("batch_id") != batch or response.get("as_of") != manifest["as_of"]:
            raise DataError("EXPORT_BATCH_MISMATCH")
        current = response["period"]
        if period is not None and current != period:
            raise DataError("EXPORT_PERIOD_MISMATCH")
        period = current
        rows = response["rows"]
        if any(row.get("taxonomy") != taxonomy or row.get("level") != level for row in rows):
            raise DataError("EXPORT_CLASSIFICATION_MISMATCH")
        return rows

    tables = {}
    if fmt == "csv":
        tax = params.get("taxonomy", "SW")
        level = params.get("level_or_series", "L1")
        rows = filter_and_sort(fetch(tax, level), params)
        # Only sortable columns present in the current page may be explicit.
        allowed = {"name", "status", *PAGE_KEYS[params["page"]]}
        if params.get("sort_direction", "default") != "default" and params.get("sort_key") not in allowed:
            raise DataError("INVALID_SORT_KEY")
        rows_by_tax[tax] = rows
        tables["当前视图"] = _industry_table(rows, batch, period, PAGE_KEYS[params["page"]], publication_state=manifest.get("publication_state"))
    else:
        for tax, sheet_name in zip(TAXONOMIES, SHEET_NAMES[:4]):
            rows = []
            for level in tax["levels"]:
                rows.extend(sort_rows(fetch(tax["id"], level["id"]), page="moneyflow"))
            rows_by_tax[tax["id"]] = rows
            tables[sheet_name] = _industry_table(rows, batch, period, tuple(METRICS) + tuple(COUNTS), publication_state=manifest.get("publication_state"), include_rank=True)
        tables["口径说明"] = _notes(batch, manifest, period, params)
        tables["审计"] = _audit(query, manifest, period, rows_by_tax)
    all_uids = [row["uid"] for rows in rows_by_tax.values() for row in rows]
    if len(set(all_uids)) != len(all_uids):
        raise DataError("EXPORT_DUPLICATE_INDUSTRY")
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, filename = tempfile.mkstemp(prefix=".export-", suffix="." + fmt, dir=destination.parent)
    os.close(fd)
    temporary = Path(filename)
    try:
        audit = _write_csv(temporary, tables["当前视图"]) if fmt == "csv" else _write_xlsx(temporary, tables)
        with temporary.open("rb") as stream:
            os.fsync(stream.fileno())
        content = temporary.read_bytes()
        # Atomic create refuses a concurrent collision instead of replacing a
        # previously exported user's file. The staging file is always removed.
        try:
            os.link(temporary, destination)
        except FileExistsError:
            raise DataError("EXPORT_DESTINATION_EXISTS") from None
        directory = os.open(destination.parent, os.O_RDONLY)
        try:
            os.fsync(directory)
        finally:
            os.close(directory)
        return {"path": str(destination.resolve()), "filename": destination.name, "format": fmt, "batch_id": batch,
                "period": period, "row_count": len(all_uids), "sha256": hashlib.sha256(content).hexdigest(),
                "bytes": len(content), "audit": audit}
    finally:
        temporary.unlink(missing_ok=True)
